"""Cópia auditável e carregamento da base original.

A base original nunca é alterada: ela é copiada para data/raw/ e os hashes
SHA-256 da origem e da cópia são comparados e registrados.
"""
from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import openpyxl
import pandas as pd

from src.utils.paths import load_config, raw_copy_path, raw_source_path

EXPECTED_COLUMNS: list[str] = [
    "Timestamp", "B1", "B2", "B3", "LT01", "FT01", "FT02", "FT03", "PT01", "PT02",
]


def sha256(path: Path, chunk_size: int = 1 << 20) -> str:
    """Hash SHA-256 de um arquivo, lido em blocos."""
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def copy_raw(source: Path | None = None, dest: Path | None = None) -> dict[str, str]:
    """Copia a base original para data/raw/ e valida a integridade.

    Se a cópia já existir com o mesmo hash, ela é reutilizada. Se existir com
    hash diferente, levanta erro (nunca sobrescreve silenciosamente).
    """
    source = source or raw_source_path()
    dest = dest or raw_copy_path()
    if not source.is_file():
        raise FileNotFoundError(f"Base original não encontrada: {source}")

    source_hash = sha256(source)
    if dest.is_file():
        dest_hash = sha256(dest)
        if dest_hash != source_hash:
            raise RuntimeError(
                f"A cópia em {dest} difere da origem (hash {dest_hash[:12]} != "
                f"{source_hash[:12]}). Verifique antes de substituir.")
        action = "reutilizada"
    else:
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, dest)
        dest_hash = sha256(dest)
        if dest_hash != source_hash:
            raise RuntimeError("Hash da cópia difere da origem logo após a cópia.")
        action = "criada"

    return {"source": str(source), "dest": str(dest), "sha256_source": source_hash,
            "sha256_dest": dest_hash, "action": action}


def read_raw_excel(path: Path | None = None) -> tuple[pd.DataFrame, dict[str, int]]:
    """Lê a aba de dados sem alterar valores.

    Remove apenas linhas e colunas totalmente vazias (resíduo de formatação do
    Excel) e devolve a contagem do que foi removido para registro.
    """
    cfg = load_config()
    path = path or raw_copy_path()
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        ws = wb[cfg["paths"]["raw_sheet"]]
        used_rows, used_cols = ws.max_row, ws.max_column
    finally:
        wb.close()
    df = pd.read_excel(path, sheet_name=cfg["paths"]["raw_sheet"], engine="openpyxl")
    shape_read = df.shape

    empty_cols = [c for c in df.columns if df[c].isna().all()]
    df = df.drop(columns=empty_cols)
    n_empty_rows = int(df.isna().all(axis=1).sum())
    df = df.loc[~df.isna().all(axis=1)].reset_index(drop=True)

    missing = [c for c in EXPECTED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"Colunas esperadas ausentes: {missing}")

    ts_col = cfg["data"]["timestamp_column"]
    df[ts_col] = pd.to_datetime(df[ts_col], errors="raise")

    info = {"excel_used_rows": used_rows, "excel_used_cols": used_cols,
            "rows_read": shape_read[0], "cols_read": shape_read[1],
            "empty_rows_removed": n_empty_rows, "empty_cols_removed": len(empty_cols)}
    return df, info
